from pathlib import Path
import pandas as pd
from openpyxl import load_workbook


def read_small_excel(path: Path) -> pd.DataFrame:
    return pd.read_excel(path, engine="openpyxl")


def iter_excel_batches(path: Path, batch_size: int = 20_000):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows))
    batch = []
    try:
        for row in rows:
            batch.append(row)
            if len(batch) >= batch_size:
                yield pd.DataFrame(batch, columns=headers)
                batch = []
        if batch:
            yield pd.DataFrame(batch, columns=headers)
    finally:
        wb.close()
